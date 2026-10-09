"""
PakTradeX — 6-Slide Pitch Deck for AI National Hackathon
- Exactly 6 slides
- Follows: Problem -> Solution -> Impact -> Innovation -> Proof
- NO tech jargon (no Flutter, FastAPI, Docker, SQL, etc.)
- Clean visual diagrams (process workflows, arrows, comparison cards)
- PowerPoint slide transitions enabled (smooth fade)
- Simple, easy-to-understand words for judges
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Clean, elegant color palette
    BG_CANVAS     = RGBColor(248, 250, 252)  # Slate 50
    BG_CARD       = RGBColor(255, 255, 255)  # Pure White
    COLOR_TITLE   = RGBColor(15, 23, 42)     # Slate 900
    COLOR_BODY    = RGBColor(51, 65, 85)     # Slate 700
    COLOR_MUTED   = RGBColor(100, 116, 139)  # Slate 500
    COLOR_GREEN   = RGBColor(16, 149, 106)   # PSX Green
    COLOR_GREEN_LT= RGBColor(209, 250, 229)  # Green tint
    COLOR_BLUE    = RGBColor(2, 132, 199)    # Sky 600
    COLOR_BLUE_LT = RGBColor(224, 242, 254)  # Blue tint
    COLOR_RED     = RGBColor(225, 29, 72)    # Rose 600
    COLOR_RED_LT  = RGBColor(255, 228, 230)  # Rose tint
    COLOR_GOLD    = RGBColor(217, 119, 6)    # Amber 600
    COLOR_GOLD_LT = RGBColor(254, 243, 199)  # Amber tint
    COLOR_BORDER  = RGBColor(226, 232, 240)  # Slate 200

    def apply_transition(slide):
        """Adds smooth fade transition to slide."""
        try:
            tr = parse_xml('<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:fade/></p:transition>')
            slide._element.append(tr)
        except Exception:
            pass

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_CANVAS
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, sub_text, tag_col=COLOR_GREEN, tag_bg=COLOR_GREEN_LT):
        # Tag pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(2.8), Inches(0.36))
        pill.fill.solid()
        pill.fill.fore_color.rgb = tag_bg
        pill.line.color.rgb = tag_col
        tf_pill = pill.text_frame
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = tag_text.upper()
        p_pill.font.size = Pt(10.5)
        p_pill.font.bold = True
        p_pill.font.color.rgb = tag_col
        p_pill.alignment = PP_ALIGN.CENTER

        # Title & Subtitle Box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TITLE

        p_s = tf.add_paragraph()
        p_s.text = sub_text
        p_s.font.size = Pt(13.5)
        p_s.font.color.rgb = COLOR_MUTED
        p_s.space_before = Pt(3)

    # ─────────────────────────────────────────────────────────────
    # SLIDE 1: Title & Framework Roadmap
    # ─────────────────────────────────────────────────────────────
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_transition(s1)
    add_bg(s1)

    # Main Card
    main_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    main_card.fill.solid()
    main_card.fill.fore_color.rgb = BG_CARD
    main_card.line.color.rgb = COLOR_BORDER

    tb1 = s1.shapes.add_textbox(Inches(1.3), Inches(1.3), Inches(10.7), Inches(3.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "AI NATIONAL HACKATHON PAKISTAN  •  CITY UNIVERSITY"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    p = tf1.add_paragraph()
    p.text = "PakTradeX"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    p.space_before = Pt(6)

    p = tf1.add_paragraph()
    p.text = "Pakistan Stock Exchange (PSX) Mobile Trading App"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE
    p.space_before = Pt(4)

    p = tf1.add_paragraph()
    p.text = "A simple, risk-free mobile app that helps everyday Pakistanis learn, practice, and trade stocks using real market prices and AI guidance."
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_BODY
    p.space_before = Pt(10)

    # Visual 5-Stage Roadmap Diagram
    steps_roadmap = ["1. Problem", "2. Solution", "3. Impact", "4. Innovation", "5. Proof"]
    box_w = 2.0
    for i, step in enumerate(steps_roadmap):
        x = Inches(1.3 + (i * 2.15))
        y = Inches(4.8)
        box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(box_w), Inches(0.8))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_GREEN_LT if i == 0 else BG_CANVAS
        box.line.color.rgb = COLOR_GREEN if i == 0 else COLOR_BORDER

        tf_b = box.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = step
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_GREEN if i == 0 else COLOR_TITLE
        p_b.alignment = PP_ALIGN.CENTER

        # Arrow between steps
        if i < 4:
            arrow = s1.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(box_w + 0.03), Inches(5.1), Inches(0.09), Inches(0.18))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = COLOR_MUTED
            arrow.line.fill.background()

    # ─────────────────────────────────────────────────────────────
    # SLIDE 2: 1. The Problem (Visual 3-Barrier Diagram)
    # ─────────────────────────────────────────────────────────────
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_transition(s2)
    add_bg(s2)
    add_header(s2, "Stage 1", "The Problem & Who It Affects", "Why less than 0.1% of Pakistan's population invests in the stock market.", COLOR_RED, COLOR_RED_LT)

    problems = [
        ("Barrier 1: Fear of Loss", "No Way to Practice First", "People want to invest, but they are terrified of losing their hard-earned money. There is no simple place to learn without real risk.", COLOR_RED),
        ("Barrier 2: Confusing Apps", "Complicated Terms & Jargon", "Existing broker apps look like 1990s accounting software. They are slow, confusing, and full of complex words beginners don't understand.", COLOR_GOLD),
        ("Barrier 3: Hard Onboarding", "Heavy Paperwork & Delays", "Opening a real broker account takes days, physical forms, and long verification steps, turning away young mobile users.", COLOR_BLUE),
    ]

    for i, (title, sub, desc, col) in enumerate(problems):
        x = Inches(0.8 + (i * 3.98))
        box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.2), Inches(3.78), Inches(3.8))
        box.fill.solid()
        box.fill.fore_color.rgb = BG_CARD
        box.line.color.rgb = col
        box.line.width = Pt(1.5)

        tb = s2.shapes.add_textbox(x + Inches(0.2), Inches(2.4), Inches(3.38), Inches(3.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"0{i+1}."
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE
        p.space_before = Pt(4)

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_before = Pt(2)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_BODY
        p.space_before = Pt(10)

    # Bottom Target Persona Banner
    target_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.733), Inches(0.65))
    target_box.fill.solid()
    target_box.fill.fore_color.rgb = COLOR_RED_LT
    target_box.line.color.rgb = COLOR_RED
    tf_t = target_box.text_frame
    pt = tf_t.paragraphs[0]
    pt.text = "🎯 Who It Affects: College students, young job holders, and first-time savers across Pakistan who want to grow their savings."
    pt.font.size = Pt(12)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_RED
    pt.alignment = PP_ALIGN.CENTER

    # ─────────────────────────────────────────────────────────────
    # SLIDE 3: 2. Solution (Visual 4-Step User Journey Diagram)
    # ─────────────────────────────────────────────────────────────
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_transition(s3)
    add_bg(s3)
    add_header(s3, "Stage 2", "Our Solution & The User Journey", "A friendly, risk-free mobile app built for beginners.", COLOR_GREEN, COLOR_GREEN_LT)

    sol_steps = [
        ("Step 1", "Practice Free", "Users get 1,000,000 PKR virtual balance to practice trading with zero risk."),
        ("Step 2", "Live PSX Prices", "Streams genuine live prices and charts for top companies (MCB, UBL, Meezan, OGDC)."),
        ("Step 3", "Halal Stock Filter", "Easily identify Islamic Shariah-compliant halal stocks and track KMI-30."),
        ("Step 4", "Easy Wallets", "Practice 1-click verification and simulated deposits using Raast, JazzCash & EasyPaisa."),
    ]

    for i, (step_num, title, desc) in enumerate(sol_steps):
        x = Inches(0.8 + (i * 2.98))
        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.2), Inches(2.78), Inches(3.8))
        box.fill.solid()
        box.fill.fore_color.rgb = BG_CARD
        box.line.color.rgb = COLOR_GREEN
        box.line.width = Pt(1.5)

        tb = s3.shapes.add_textbox(x + Inches(0.2), Inches(2.4), Inches(2.38), Inches(3.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = step_num.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_GREEN

        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE
        p.space_before = Pt(4)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_BODY
        p.space_before = Pt(10)

        # Arrow to next step
        if i < 3:
            arrow = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(2.78 + 0.05), Inches(4.0), Inches(0.1), Inches(0.2))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = COLOR_GREEN
            arrow.line.fill.background()

    # Bottom audience banner
    bar3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.733), Inches(0.65))
    bar3.fill.solid()
    bar3.fill.fore_color.rgb = COLOR_GREEN_LT
    bar3.line.color.rgb = COLOR_GREEN
    tf_bar3 = bar3.text_frame
    pb3 = tf_bar3.paragraphs[0]
    pb3.text = "👥 Target Audience: Beginners, smartphone users, and students taking their first steps into financial independence."
    pb3.font.size = Pt(12)
    pb3.font.bold = True
    pb3.font.color.rgb = COLOR_GREEN
    pb3.alignment = PP_ALIGN.CENTER

    # ─────────────────────────────────────────────────────────────
    # SLIDE 4: 3. Impact (Visual 3-Pillar Diagram)
    # ─────────────────────────────────────────────────────────────
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_transition(s4)
    add_bg(s4)
    add_header(s4, "Stage 3", "Practical Impact & National Need", "Promoting financial literacy and helping people beat inflation.", COLOR_BLUE, COLOR_BLUE_LT)

    impact_pillars = [
        ("Pillar 1: Beating Inflation", "Protecting Savings", "Cash in bank accounts loses value due to inflation. PakTradeX teaches citizens how to invest in strong, dividend-paying companies to protect their wealth.", COLOR_BLUE),
        ("Pillar 2: Learning by Doing", "Practical Skills", "Reading books or listening to lectures is boring. By placing real practice trades on live moving charts, users build true trading confidence.", COLOR_GREEN),
        ("Pillar 3: Growing the Economy", "Fueling National Growth", "By training thousands of confident new investors, we create a feeder pipeline for the stock exchange, bringing domestic capital into Pakistani businesses.", COLOR_GOLD),
    ]

    for i, (title, sub, desc, col) in enumerate(impact_pillars):
        x = Inches(0.8 + (i * 3.98))
        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.2), Inches(3.78), Inches(4.6))
        box.fill.solid()
        box.fill.fore_color.rgb = BG_CARD
        box.line.color.rgb = col
        box.line.width = Pt(1.5)

        tb = s4.shapes.add_textbox(x + Inches(0.2), Inches(2.4), Inches(3.38), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"★ {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE
        p.space_before = Pt(4)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_BODY
        p.space_before = Pt(12)

    # ─────────────────────────────────────────────────────────────
    # SLIDE 5: 4. Innovation (Visual Side-by-Side Comparison)
    # ─────────────────────────────────────────────────────────────
    s5 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_transition(s5)
    add_bg(s5)
    add_header(s5, "Stage 4", "The Innovation & Why We Use AI", "Solving real user frustration without complex buzzwords.", COLOR_GOLD, COLOR_GOLD_LT)

    # Left Card: Traditional Way
    box_old = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(5.7), Inches(4.6))
    box_old.fill.solid()
    box_old.fill.fore_color.rgb = COLOR_RED_LT
    box_old.line.color.rgb = COLOR_RED
    box_old.line.width = Pt(1.5)

    tb_old = s5.shapes.add_textbox(Inches(1.1), Inches(2.4), Inches(5.1), Inches(4.2))
    tf_old = tb_old.text_frame
    tf_old.word_wrap = True

    p = tf_old.paragraphs[0]
    p.text = "❌ The Old Way (The Barrier)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED

    pts_old = [
        "50-page complex PDF financial reports published quarterly.",
        "Confusing terms (P/E ratio, Dividend Yield, EBITDA) left unexplained.",
        "English-only documents that regular people cannot understand.",
        "Zero instant help: users rely on unverified rumors on social media.",
    ]
    for pt in pts_old:
        p = tf_old.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TITLE
        p.space_before = Pt(10)

    # Right Card: PakTradeX AI Way
    box_new = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.6))
    box_new.fill.solid()
    box_new.fill.fore_color.rgb = COLOR_GREEN_LT
    box_new.line.color.rgb = COLOR_GREEN
    box_new.line.width = Pt(1.5)

    tb_new = s5.shapes.add_textbox(Inches(7.1), Inches(2.4), Inches(5.1), Inches(4.2))
    tf_new = tb_new.text_frame
    tf_new.word_wrap = True

    p = tf_new.paragraphs[0]
    p.text = "✅ The PakTradeX AI Way (The Innovation)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    pts_new = [
        "3-Bullet Simple Summary explaining if a company is healthy and growing.",
        "Answers stock questions in plain everyday Urdu and English.",
        "Live Risk Warnings: Alerts beginners if a stock is too volatile or risky.",
        "Zero-Downtime Reliability: Works seamlessly online and offline with zero crashes.",
    ]
    for pt in pts_new:
        p = tf_new.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TITLE
        p.space_before = Pt(10)

    # ─────────────────────────────────────────────────────────────
    # SLIDE 6: 5. Proof (Visual 4-Card Evidence Grid)
    # ─────────────────────────────────────────────────────────────
    s6 = prs.slides.add_slide(prs.slide_layouts[6])
    apply_transition(s6)
    add_bg(s6)
    add_header(s6, "Stage 5", "Proof of Implementation & Working Prototype", "A completed, fully tested, and working product ready right now.", COLOR_GREEN, COLOR_GREEN_LT)

    proofs = [
        ("Working Mobile App (APK)", "Installed on Mobile", "Full 53.6 MB Android release app running smoothly on physical smartphones with fast navigation.", COLOR_GREEN),
        ("31 Real PSX Stocks", "Real Live Prices", "Authentic prices and charts for top companies (MCB, UBL, Meezan Bank, Systems Ltd, OGDC).", COLOR_BLUE),
        ("Full Trading Engine", "Buy & Sell Execution", "Place Market and Limit orders with live profit, loss, and portfolio balance updates.", COLOR_GOLD),
        ("Tested & Reliable", "18/18 Tests Passed", "Every feature has been tested with automated test suites, and all code is published on GitHub.", COLOR_RED),
    ]

    for i, (title, sub, desc, col) in enumerate(proofs):
        r = i // 2
        c = i % 2
        x = Inches(0.8 + (c * 5.95))
        y = Inches(2.2 + (r * 2.3))

        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.1))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        tb = s6.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), Inches(5.35), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"✔ {title}"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE
        p.space_before = Pt(2)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_BODY
        p.space_before = Pt(6)

    bar6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.6), Inches(11.733), Inches(0.55))
    bar6.fill.solid()
    bar6.fill.fore_color.rgb = BG_CANVAS
    bar6.line.color.rgb = COLOR_BORDER
    tf_bar6 = bar6.text_frame
    pb6 = tf_bar6.paragraphs[0]
    pb6.text = "🎯 Ready for Live Demonstration: Working prototype available for live judge testing."
    pb6.font.size = Pt(11)
    pb6.font.bold = True
    pb6.font.color.rgb = COLOR_GREEN
    pb6.alignment = PP_ALIGN.CENTER

    output_path = "PakTradeX_Hackathon_Final_Deck.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
